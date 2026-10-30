"""Tiny helpers around python-docx for a journal-style report."""
from __future__ import annotations

import re

from docx import Document
from docx.oxml import parse_xml
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


class Report:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        for s in self.doc.sections:
            s.left_margin = s.right_margin = Cm(2.0)
            s.top_margin = s.bottom_margin = Cm(2.0)
        for lvl, size in ((1, 13), (2, 11.5), (3, 10.5)):
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.name = "Times New Roman"
            h.font.size = Pt(size)
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        self.n_fig = 0
        self.n_tab = 0
        self.n_eq = 0

    # ------------------------------------------------------------------ text
    def title(self, text, sub=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(16)
        if sub:
            for line in sub:
                q = self.doc.add_paragraph()
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                q.add_run(line).font.size = Pt(10.5)

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text, italic=False, bold_lead=None, align="justify", size=None):
        """Paragraph. Supports **bold** and *italic* inline markers."""
        par = self.doc.add_paragraph()
        par.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                         "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        par.paragraph_format.space_after = Pt(4)
        if bold_lead:
            r = par.add_run(bold_lead + " ")
            r.bold = True
        self._inline(par, text, italic, size)
        return par

    def _inline(self, par, text, italic=False, size=None, bold=False):
        for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            if not tok:
                continue
            if tok.startswith("**"):
                _script_runs(par, tok[2:-2], bold=True, italic=italic, size=size)
            elif tok.startswith("*"):
                _script_runs(par, tok[1:-1], bold=bold, italic=True, size=size)
            else:
                _script_runs(par, tok, bold=bold, italic=italic, size=size)

    def bullets(self, items, numbered=False):
        for it in items:
            par = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            par.paragraph_format.space_after = Pt(2)
            self._inline(par, it)

    def eq(self, latex, fallback, label=True):
        """Display equation, right-numbered. Written as a native Word equation (OMML) converted from
        LaTeX; if Office's MathML->OMML stylesheet is not available, the Unicode `fallback` is used."""
        self.n_eq += 1
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t.rows[0].cells
        c0.width, c1.width = Cm(15), Cm(2)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        omml = _latex_to_omml(latex)
        if omml is not None:
            p0._p.append(omml)
        else:
            r = p0.add_run(fallback)
            r.font.name = "Cambria Math"
            r.italic = True
            r.font.size = Pt(10.5)
        if label:
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p1.add_run(f"({self.n_eq})")
        return self.n_eq

    # ------------------------------------------------------------------ floats
    def figure(self, path, caption, width_cm=16.5):
        self.n_fig += 1
        self.doc.add_picture(str(path), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = c.add_run(f"Fig. {self.n_fig}. ")
        r.bold = True
        r.font.size = Pt(9)
        self._inline(c, caption, size=9)
        return self.n_fig

    def table(self, header, rows, caption, col_widths=None, font=8.5, bold_rows=()):
        self.n_tab += 1
        c = self.doc.add_paragraph()
        r = c.add_run(f"Table {self.n_tab}. ")
        r.bold = True
        r.font.size = Pt(9)
        self._inline(c, caption, size=9)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, hd in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.text = ""
            _script_runs(cell.paragraphs[0], str(hd), bold=True, size=font)
            _shade(cell, "D9D9D9")
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                cell = t.rows[i].cells[j]
                cell.text = ""
                _script_runs(cell.paragraphs[0], str(v), bold=i - 1 in bold_rows, size=font)
        if col_widths:
            for row in t.rows:
                for j, wdt in enumerate(col_widths):
                    row.cells[j].width = Cm(wdt)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return self.n_tab

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(str(path))


# x_k, x_{ab}, x^2, e^{-jw}: a word followed by a sub/superscript marker
_SCRIPT = re.compile(r"([_^])(\{[^{}]*\}|[\w\u0300-\u036f]+)")
# tokens that are code identifiers or file names, not math
_IDENT = re.compile(r"\.py|\.json|\.csv|\.txt|\.docx|`|^[a-z]{2,}_[A-Za-z]|^\d+_|^[\w.-]+/[\w./*<>-]+$")


def _script_runs(par, text, bold=False, italic=False, size=None):
    """Add `text` to `par`, rendering math-style x_sub / x^{sup} as real subscript/superscript runs."""
    parts = []                                   # (text, 'sub' | 'sup' | None)
    for word in re.split(r"(\s+)", text):
        if not word or word.isspace() or _IDENT.search(word):
            parts.append((word, None))
            continue
        pos = 0
        for m in _SCRIPT.finditer(word):
            if m.start() == 0:                   # nothing to attach the script to
                continue
            parts.append((word[pos:m.start()], None))
            body = m.group(2)
            parts.append((body[1:-1] if body.startswith("{") else body, "sub" if m.group(1) == "_" else "sup"))
            pos = m.end()
        parts.append((word[pos:], None))
    merged = []
    for s, kind in parts:
        if s and merged and merged[-1][1] == kind:
            merged[-1] = (merged[-1][0] + s, kind)
        elif s:
            merged.append((s, kind))
    for s, kind in merged:
        r = par.add_run(s)
        r.bold, r.italic = bold or None, italic or None
        if kind == "sub":
            r.font.subscript = True
        elif kind == "sup":
            r.font.superscript = True
        if size:
            r.font.size = Pt(size)


def _find_mml2omml():
    import glob
    import os
    for root in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                 os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
        hits = glob.glob(os.path.join(root, "Microsoft Office", "**", "MML2OMML.XSL"), recursive=True)
        if hits:
            return hits[0]
    return None


_XSLT = []


def _latex_to_omml(latex):
    """LaTeX -> MathML (latex2mathml) -> OMML (Office's MML2OMML.XSL). None if either is unavailable."""
    try:
        from latex2mathml.converter import convert
        from lxml import etree
    except ImportError:
        return None
    if not _XSLT:
        path = _find_mml2omml()
        _XSLT.append(etree.XSLT(etree.parse(path)) if path else None)
    if _XSLT[0] is None:
        return None
    # MML2OMML ignores <mspace>; keep \, \; \quad \qquad as real space characters
    def _space(m):
        w = float(m.group(1))
        return "<mtext>" + (" " * round(w) if w >= 1 else " ") + "</mtext>"
    mml = re.sub(r'<mspace width="([\d.]+)em"\s*/>', _space, convert(latex))
    omml = _XSLT[0](etree.fromstring(mml))
    return parse_xml(etree.tostring(omml.getroot()))


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)
